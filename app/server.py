# Copyright (c) 2025 Oleksandr Niemchenko. All Rights Reserved.
import csv
import io
import json
import logging
import os
import re
import subprocess
import tempfile
import threading
import time
import urllib.parse
import urllib.request
import uuid
from pathlib import Path
from typing import Iterable

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from pydantic import BaseModel
from openai import OpenAI
from dotenv import load_dotenv
from fastapi.staticfiles import StaticFiles
from fastapi.responses import RedirectResponse

from . import db

from .asr_whisper import transcribe_with_fallback

LOGGER = logging.getLogger(__name__)
PROMPT_PATH = Path(__file__).with_name("prompts") / "system_hr.md"
FALLBACK_SYSTEM_PROMPT = (
    "Ты — локальный HR-ассистент. Работай строго по фактам, не выдумывай данные и"
    " отвечай кратко на русском языке. Если информации не хватает, прямо укажи на"
    " это и предложи следующий шаг."
)
PROMPT_LANGUAGE_SUFFIX = (
    "\n\nТребование: отвечай ТОЛЬКО на русском языке. Не используй другие языки."
)

load_dotenv()

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OLLAMA_MODELS_DIR = (
    os.getenv("OLLAMA_MODELS_DIR")
    or os.getenv("OLLAMA_MODELS")
    or str(PROJECT_ROOT / "ollama_models" / "models")
)
os.environ["OLLAMA_MODELS"] = OLLAMA_MODELS_DIR
os.environ["OLLAMA_MODELS_DIR"] = OLLAMA_MODELS_DIR
Path(OLLAMA_MODELS_DIR).mkdir(parents=True, exist_ok=True)
OLLAMA_ENV = {**os.environ, "OLLAMA_MODELS": OLLAMA_MODELS_DIR}

ENV = os.getenv("ENV", "local")
BASE_URL = os.getenv("LOCAL_BASE_URL", "http://localhost:11434/v1")
API_KEY  = os.getenv("LOCAL_API_KEY", "LOCAL-KEY")
MODEL    = os.getenv("LOCAL_MODEL", "qwen2.5:3b")
TEMP     = float(os.getenv("TEMP", "0.25"))
MAX_TOK  = int(os.getenv("MAX_TOKENS", "512"))
NUM_CTX = int(os.getenv("NUM_CTX", "32768"))
ADMIN_TOKEN = os.getenv("ADMIN_TOKEN", "")
MEM_TURNS = int(os.getenv("MEM_TURNS", "6"))  # сколько пар (user+assistant) храним в памяти на сессию
UI_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "web_ui"))

MODEL_RECOMMENDATIONS = [
    {"name": "qwen2.5:3b", "label": "Лёгкая (≈3B)"},
    {"name": "qwen2.5:7b-instruct", "label": "Инструкционная 7B"},
    {"name": "qwen2.5-7b-long", "label": "Длинный контекст 7B"},
    {"name": "llama3:8b", "label": "LLaMA 3 8B"},
    {"name": "mistral:7b", "label": "Mistral 7B"},
]
DEFAULT_LLAMA_MODEL = os.getenv("DEFAULT_LOCAL_MODEL", MODEL_RECOMMENDATIONS[0]["name"])

MODEL_PULLS: dict[str, dict] = {}
MODEL_PULL_PROCS: dict[str, subprocess.Popen] = {}
MODEL_PULL_LOCK = threading.Lock()
AUTO_PULL_TRIGGERED = False
DEFAULT_SETTINGS = {
    "temperature": TEMP,
    "max_tokens": MAX_TOK,
    "num_ctx": NUM_CTX,
    "mem_turns": MEM_TURNS,
}

db.init_db()

client = OpenAI(base_url=BASE_URL, api_key=API_KEY)
app = FastAPI(title="LocalGPT HR API")

class ChatReq(BaseModel):
    query: str
    context: str | None = None
    session_id: str | None = None
    profile_id: int | None = None
    transcript: str | None = None
    criteria: list[str] | None = None
    model: str | None = None
    model_name: str | None = None
    title: str | None = None


class ModelPullRequest(BaseModel):
    model: str


class ProfileCreate(BaseModel):
    name: str
    description: str | None = None
    color: str | None = None


class SessionCreate(BaseModel):
    profile_id: int
    title: str | None = None
    model_name: str | None = None


class ModelSettingsRequest(BaseModel):
    profile_id: int
    model_name: str
    temperature: float | None = None
    max_tokens: int | None = None
    num_ctx: int | None = None
    mem_turns: int | None = None


class ModelSettingsReset(BaseModel):
    profile_id: int
    model_name: str


class TranscriptCreate(BaseModel):
    profile_id: int
    session_id: str | None = None
    name: str
    content: str
    source_path: str | None = None
    language: str | None = None
    timeline: list[dict] | None = None
    role_keywords: dict | None = None
    labels: dict | None = None


class TranscriptUpdate(BaseModel):
    name: str | None = None
    content: str | None = None
    language: str | None = None
    timeline: list[dict] | None = None
    role_keywords: dict | None = None
    labels: dict | None = None

def _system_prompt() -> str:
    try:
        with PROMPT_PATH.open("r", encoding="utf-8") as f:
            base_prompt = f.read()
    except (OSError, UnicodeDecodeError) as exc:
        LOGGER.warning("Failed to read system prompt from %s: %s", PROMPT_PATH, exc)
        base_prompt = FALLBACK_SYSTEM_PROMPT
    return base_prompt + PROMPT_LANGUAGE_SUFFIX


def _get_profile_or_404(profile_id: int) -> dict:
    profile = db.get_profile(profile_id)
    if not profile or not profile.get("is_active", 1):
        raise HTTPException(status_code=404, detail="profile_not_found")
    return profile


def _effective_model_settings(profile_id: int, model_name: str) -> dict:
    return db.effective_settings(profile_id, model_name, DEFAULT_SETTINGS)


def _normalize_model_name(model_name: str | None) -> str:
    return (model_name or MODEL).strip() or MODEL


def _session_title_hint(raw_title: str | None, query: str | None) -> str:
    if raw_title and raw_title.strip():
        return raw_title.strip()
    prompt = (query or "").strip()
    if prompt:
        return (prompt[:60] + ("…" if len(prompt) > 60 else "")).strip()
    return "Новая сессия"

# --- Критерии и источники данных -------------------------------------------------

MAX_CRITERIA = 200

def _normalize_criterion(raw: str) -> str:
    if not raw:
        return ""
    text = re.sub(r"[\u200b\u200c\u200d]", "", raw)
    text = text.strip()
    text = re.sub(r"^[-•*\d\)\].\s]+", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def _criteria_from_lines(lines: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for line in lines:
        cleaned = _normalize_criterion(line)
        if not cleaned:
            continue
        key = cleaned.lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(cleaned)
        if len(result) >= MAX_CRITERIA:
            break
    return result

def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")

def _parse_csv(text: str) -> list[str]:
    stream = io.StringIO(text)
    reader = csv.reader(stream)
    lines: list[str] = []
    for row in reader:
        if not row:
            continue
        for cell in row:
            if cell and cell.strip():
                lines.append(cell)
                break
    return _criteria_from_lines(lines)

def _parse_xlsx(data: bytes) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - зависит от окружения
        raise RuntimeError("openpyxl_not_available") from exc

    workbook = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    lines: list[str] = []
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        for cell in row:
            if isinstance(cell, str) and cell.strip():
                lines.append(cell)
                break
            if isinstance(cell, (int, float)):
                lines.append(str(cell))
                break
        if len(lines) >= MAX_CRITERIA:
            break
    return _criteria_from_lines(lines)

def _parse_text_block(text: str) -> list[str]:
    if not text:
        return []
    raw_lines = re.split(r"[\r\n]+", text)
    return _criteria_from_lines(raw_lines)

def _extract_sheet_id(raw: str) -> tuple[str, str | None]:
    value = (raw or "").strip()
    if not value:
        raise ValueError("google_sheet_id_required")
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", value)
    gid = None
    if match:
        sheet_id = match.group(1)
        gid_match = re.search(r"gid=([0-9]+)", value)
        if gid_match:
            gid = gid_match.group(1)
        return sheet_id, gid
    return value, None

def _load_google_sheet_csv(sheet_id: str, gid: str | None = None) -> str:
    if not sheet_id:
        raise ValueError("google_sheet_id_required")
    params = {"format": "csv"}
    if gid:
        params["gid"] = gid
    query = urllib.parse.urlencode(params)
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?{query}"
    with urllib.request.urlopen(url, timeout=15) as response:  # type: ignore[call-arg]
        data = response.read()
    return _decode_bytes(data)

# Helper: удаляем повторы и случайные куски на других языках
def _sanitize(text: str) -> str:
    """Убираем явные повторы и случайные куски на других языках."""
    import re

    # удаляем явные китайские блоки/фразы
    text = re.sub(r"[\u4e00-\u9fff]{2,}.*?(?:[.!?。！？]|$)", "", text)

    # вырезаем служебные секции-инструкции, которые модель иногда копирует в ответ
    def _drop_section(src: str, heading: str) -> str:
        pattern = re.compile(rf"(^##\s*{heading}.*?)(?=^##\s|\Z)", re.MULTILINE | re.DOTALL)
        return re.sub(pattern, "", src)

    for meta_heading in (
        "Минимальный набор фактов",
        "Чеклист перед отправкой",
        "Checklist before sending",
    ):
        text = _drop_section(text, meta_heading)

    # удаляем строки с типовыми подсказками вроде «Теперь вы можете…»
    text = re.sub(r"^.*Теперь вы можете.*$", "", text, flags=re.MULTILINE)

    lines: list[str] = []
    seen_headings: set[str] = set()
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            if lines and lines[-1] == "":
                continue
            lines.append("")
            continue

        if stripped.startswith("### "):
            key = stripped.lower()
            if key in seen_headings:
                continue
            seen_headings.add(key)

        if lines and lines[-1].strip() == stripped:
            continue

        lines.append(line)

    return "\n".join(lines).strip()


def _list_local_models() -> list[dict]:
    cmd_json = ["ollama", "list", "--json"]
    try:
        result = subprocess.run(
            cmd_json,
            capture_output=True,
            text=True,
            check=False,
            timeout=20,
            env=OLLAMA_ENV,
        )
    except FileNotFoundError as exc:  # pragma: no cover - зависит от окружения
        raise HTTPException(status_code=500, detail="Ollama не установлена или недоступна.") from exc
    except subprocess.SubprocessError as exc:  # pragma: no cover - зависит от окружения
        raise HTTPException(status_code=500, detail="Ошибка обращения к Ollama.") from exc

    stdout = result.stdout or ""
    stderr = (result.stderr or "").strip()
    fallback_plain = result.returncode != 0 and "unknown flag: --json" in stderr.lower()

    lines: list[str] = []
    if not fallback_plain:
        if result.returncode != 0:
            msg = stderr or "не удалось получить список моделей"
            raise HTTPException(status_code=500, detail={"code": "ollama_list_failed", "message": msg})
        lines = stdout.splitlines()
    else:
        # старые версии ollama без --json: парсим табличный вывод
        try:
            plain = subprocess.run(
                ["ollama", "list"],
                capture_output=True,
                text=True,
                check=False,
                timeout=20,
                env=OLLAMA_ENV,
            )
        except Exception as exc:  # pragma: no cover - зависит от окружения
            raise HTTPException(status_code=500, detail={"code": "ollama_list_failed", "message": str(exc)}) from exc
        if plain.returncode != 0:
            msg = (plain.stderr or plain.stdout or "не удалось получить список моделей").strip()
            raise HTTPException(status_code=500, detail={"code": "ollama_list_failed", "message": msg})
        lines = (plain.stdout or "").splitlines()
        # скипаем заголовок/разделители
        lines = [line for line in lines if line.strip() and not line.lower().startswith("name ")]

    models: list[dict] = []
    for line in lines:
        raw = line.strip()
        if not raw:
            continue
        try:
            payload = json.loads(raw)
            name = (payload.get("model") or payload.get("name") or "").strip()
            size = payload.get("size")
            modified_at = payload.get("modified_at")
        except json.JSONDecodeError:
            parts = raw.split()
            if not parts:
                continue
            name = parts[0]
            size = None
            modified_at = None
        models.append(
            {
                "name": name,
                "size": size,
                "size_gb": round(size / (1024**3), 2) if isinstance(size, (int, float)) else None,
                "modified_at": modified_at,
            }
        )
    return models


def _public_pull_status(data: dict) -> dict:
    copy = dict(data)
    copy.pop("_last_completed", None)
    copy.pop("_last_ts", None)
    return copy


def _pull_statuses() -> list[dict]:
    with MODEL_PULL_LOCK:
        return [_public_pull_status(item) for item in MODEL_PULLS.values()]


def _update_pull(task_id: str, **fields) -> None:
    with MODEL_PULL_LOCK:
        if task_id not in MODEL_PULLS:
            return
        MODEL_PULLS[task_id].update(fields)
        if fields.get("state") == "canceled":
            MODEL_PULL_PROCS.pop(task_id, None)


def _run_model_pull(task_id: str, model: str) -> None:
    start_ts = time.time()
    _update_pull(task_id, state="running", started_at=start_ts)

    def _run(cmd: list[str], parse_json: bool) -> tuple[int, str | None]:
        nonlocal start_ts
        last_completed = 0
        last_ts = start_ts
        raw_error = None
        try:
            process = subprocess.Popen(
                cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, env=OLLAMA_ENV
            )
            with MODEL_PULL_LOCK:
                MODEL_PULL_PROCS[task_id] = process
        except FileNotFoundError:
            _update_pull(task_id, state="error", message="Ollama не установлена", finished_at=time.time())
            return 1, "Ollama не установлена"
        except Exception as exc:  # pragma: no cover - зависит от окружения
            _update_pull(task_id, state="error", message=str(exc), finished_at=time.time())
            return 1, str(exc)

        try:
            assert process.stdout is not None
            for raw_line in process.stdout:
                line = raw_line.strip()
                if not line:
                    continue
                with MODEL_PULL_LOCK:
                    if MODEL_PULLS.get(task_id, {}).get("state") == "canceled":
                        process.terminate()
                        break
                if parse_json:
                    try:
                        payload = json.loads(line)
                    except json.JSONDecodeError:
                        if "unknown flag: --json" in line.lower():
                            raw_error = "unknown flag: --json"
                            break
                        continue

                    status_text = payload.get("status") or payload.get("digest") or ""
                    total = payload.get("total")
                    completed = payload.get("completed")

                    percent = None
                    speed_mbps = None
                    eta_seconds = None

                    if isinstance(total, (int, float)) and isinstance(completed, (int, float)) and total > 0:
                        percent = round(100 * completed / total, 2)
                        now = time.time()
                        delta_bytes = completed - last_completed
                        delta_time = max(now - last_ts, 1e-3)
                        if delta_bytes > 0:
                            speed_mbps = round(delta_bytes / delta_time / (1024 * 1024), 2)
                            remaining = total - completed
                            eta_seconds = round(remaining / (delta_bytes / delta_time)) if delta_bytes else None
                        last_completed = completed
                        last_ts = now

                    _update_pull(
                        task_id,
                        state="running",
                        message=status_text,
                        completed=completed,
                        total=total,
                        percent=percent,
                        speed_mbps=speed_mbps,
                        eta_seconds=eta_seconds,
                        _last_completed=last_completed,
                        _last_ts=last_ts,
                    )
                else:
                    _update_pull(task_id, state="running", message=line)

            process.wait()
            finished = time.time()
            with MODEL_PULL_LOCK:
                current_state = MODEL_PULLS.get(task_id, {}).get("state")
            if current_state == "canceled":
                _update_pull(task_id, state="canceled", message="Остановлено пользователем", finished_at=finished)
                return process.returncode, raw_error
            if process.returncode == 0:
                _update_pull(task_id, state="completed", message="Готово", finished_at=finished)
            else:
                _update_pull(
                    task_id,
                    state="error",
                    message=f"ollama pull вернул код {process.returncode}",
                    finished_at=finished,
                )
            return process.returncode, raw_error
        finally:
            with MODEL_PULL_LOCK:
                if task_id in MODEL_PULLS:
                    MODEL_PULLS[task_id]["finished_at"] = MODEL_PULLS[task_id].get("finished_at") or time.time()
                MODEL_PULL_PROCS.pop(task_id, None)

    rc, err = _run(["ollama", "pull", model, "--json"], parse_json=True)
    if err and "unknown flag: --json" in err.lower():
        _update_pull(task_id, state="running", message="Старая версия Ollama, включён текстовый прогресс")
        _run(["ollama", "pull", model], parse_json=False)


def _start_model_pull(model: str) -> dict:
    task_id = uuid.uuid4().hex
    entry = {
        "task_id": task_id,
        "model": model,
        "state": "pending",
        "message": "Запуск загрузки…",
        "started_at": None,
        "finished_at": None,
        "completed": None,
        "total": None,
    }
    with MODEL_PULL_LOCK:
        MODEL_PULLS[task_id] = entry

    thread = threading.Thread(target=_run_model_pull, args=(task_id, model), daemon=True)
    thread.start()
    return _public_pull_status(entry)


def _maybe_trigger_default_pull(installed: list[dict], downloads: list[dict]) -> None:
    global AUTO_PULL_TRIGGERED
    if AUTO_PULL_TRIGGERED:
        return
    if installed:
        return
    active = [item for item in downloads if item.get("state") in {"pending", "running"}]
    if active:
        return
    AUTO_PULL_TRIGGERED = True
    downloads.append(_start_model_pull(DEFAULT_LLAMA_MODEL))


def _cancel_model_pull(task_id: str) -> dict:
    with MODEL_PULL_LOCK:
        entry = MODEL_PULLS.get(task_id)
        process = MODEL_PULL_PROCS.get(task_id)
    if not entry:
        raise HTTPException(status_code=404, detail="task_not_found")
    if entry.get("state") in {"completed", "error", "canceled"}:
        return _public_pull_status(entry)

    _update_pull(task_id, state="canceled", message="Остановлено пользователем", finished_at=time.time())
    if process and process.poll() is None:
        try:
            process.terminate()
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
        finally:
            with MODEL_PULL_LOCK:
                MODEL_PULL_PROCS.pop(task_id, None)

    return _public_pull_status(MODEL_PULLS[task_id])


@app.get("/models")
def list_models():
    error = None
    try:
        installed = _list_local_models()
    except HTTPException as exc:
        installed = []
        error = exc.detail
        LOGGER.warning("Оllama list failed: %s", exc.detail)
    downloads = _pull_statuses()
    if not error:
        _maybe_trigger_default_pull(installed, downloads)
    response = {
        "installed": installed,
        "models": [item.get("name") for item in installed if item.get("name")],
        "recommended": MODEL_RECOMMENDATIONS,
        "downloads": downloads,
        "error": error,
    }
    return response


@app.post("/models/pull")
def pull_model(req: ModelPullRequest):
    model = (req.model or "").strip()
    if not model:
        raise HTTPException(status_code=422, detail={"code": "model_required", "message": "Укажите модель"})
    return _start_model_pull(model)


@app.get("/models/pull/{task_id}")
def pull_status(task_id: str):
    with MODEL_PULL_LOCK:
        entry = MODEL_PULLS.get(task_id)
    if not entry:
        raise HTTPException(status_code=404, detail="task_not_found")
    return _public_pull_status(entry)


@app.get("/models/downloads")
def pull_statuses():
    return {"downloads": _pull_statuses()}

@app.post("/models/pull/{task_id}/cancel")
def cancel_pull(task_id: str):
    return _cancel_model_pull(task_id)


@app.delete("/models/{model_name}")
def delete_model(model_name: str):
    target = (model_name or "").strip()
    if not target:
        raise HTTPException(status_code=422, detail={"code": "model_required", "message": "Укажите модель"})

    try:
        result = subprocess.run(
            ["ollama", "rm", target], capture_output=True, text=True, check=False, timeout=30, env=OLLAMA_ENV
        )
    except FileNotFoundError as exc:  # pragma: no cover - зависит от окружения
        raise HTTPException(status_code=500, detail="Ollama не установлена или недоступна.") from exc
    except subprocess.SubprocessError as exc:  # pragma: no cover - зависит от окружения
        raise HTTPException(status_code=500, detail="Ошибка удаления модели через Ollama") from exc

    if result.returncode != 0:
        msg = (result.stderr or result.stdout or "Не удалось удалить модель").strip()
        raise HTTPException(status_code=500, detail={"code": "ollama_rm_failed", "message": msg})

    return {"deleted": target}


# --- Профили и сессии -----------------------------------------------------------


@app.get("/profiles")
def list_profiles():
    profiles = db.list_profiles()
    return {"profiles": profiles}


@app.post("/profiles")
def create_profile(req: ProfileCreate):
    try:
        profile = db.create_profile(req.name, req.description, req.color)
        return profile
    except ValueError:
        raise HTTPException(status_code=422, detail={"code": "profile_name_required", "message": "Укажите имя профиля"})


@app.delete("/profiles/{profile_id}")
def delete_profile(profile_id: int):
    profiles = db.list_profiles()
    if len(profiles) <= 1:
        raise HTTPException(status_code=400, detail={"code": "cannot_delete_last_profile", "message": "Нельзя удалить единственный профиль"})
    db.delete_profile(profile_id)
    db.ensure_default_profile()
    return {"deleted": profile_id}


@app.get("/sessions")
def list_chat_sessions(profile_id: int):
    _get_profile_or_404(profile_id)
    sessions = db.list_sessions(profile_id)
    return {"sessions": sessions}


@app.post("/sessions")
def create_chat_session(req: SessionCreate):
    _get_profile_or_404(req.profile_id)
    title = _session_title_hint(req.title, None)
    model_name = _normalize_model_name(req.model_name)
    try:
        session = db.create_session(req.profile_id, title=title, model_name=model_name)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail={"code": str(exc), "message": "Укажите название сессии"}) from exc
    return session


@app.delete("/sessions/{session_id}")
def delete_chat_session(session_id: str):
    session = db.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="session_not_found")
    db.delete_session(session_id)
    return {"deleted": session_id}


@app.get("/sessions/{session_id}/messages")
def get_session_messages(session_id: str, limit: int | None = None):
    session = db.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="session_not_found")
    messages = db.get_messages(session_id, limit_pairs=limit if limit else None)
    return {"session_id": session_id, "messages": messages}


# --- Настройки модели -----------------------------------------------------------


@app.get("/model_settings")
def get_model_settings(profile_id: int, model_name: str):
    _get_profile_or_404(profile_id)
    model = _normalize_model_name(model_name)
    return _effective_model_settings(profile_id, model)


@app.post("/model_settings")
def save_model_settings(req: ModelSettingsRequest):
    _get_profile_or_404(req.profile_id)
    model = _normalize_model_name(req.model_name)
    db.upsert_model_settings(
        req.profile_id,
        model,
        temperature=req.temperature,
        max_tokens=req.max_tokens,
        num_ctx=req.num_ctx,
        mem_turns=req.mem_turns,
    )
    return _effective_model_settings(req.profile_id, model)


@app.post("/model_settings/reset_to_defaults")
def reset_model_settings(req: ModelSettingsReset):
    _get_profile_or_404(req.profile_id)
    model = _normalize_model_name(req.model_name)
    db.reset_model_settings(req.profile_id, model)
    return _effective_model_settings(req.profile_id, model)


# --- Транскрипты ---------------------------------------------------------------


@app.get("/transcripts")
def list_transcripts(profile_id: int, session_id: str | None = None):
    _get_profile_or_404(profile_id)
    transcripts = db.list_transcripts(profile_id, session_id=session_id)
    return {"transcripts": transcripts}


@app.post("/transcripts")
def create_transcript(req: TranscriptCreate):
    _get_profile_or_404(req.profile_id)
    if req.session_id:
        session = db.get_session(req.session_id)
        if not session:
            raise HTTPException(status_code=404, detail="session_not_found")
        if session["profile_id"] != req.profile_id:
            raise HTTPException(status_code=400, detail="session_profile_mismatch")
    try:
        transcript = db.create_transcript(
            profile_id=req.profile_id,
            session_id=req.session_id,
            name=req.name,
            content=req.content,
            source_path=req.source_path,
            language=req.language,
            timeline=req.timeline,
            role_keywords=req.role_keywords,
            labels=req.labels,
        )
        return transcript
    except ValueError as exc:
        raise HTTPException(status_code=422, detail={"code": str(exc), "message": "Контент транскрипта обязателен"}) from exc


@app.patch("/transcripts/{transcript_id}")
def update_transcript(transcript_id: int, req: TranscriptUpdate):
    updated = db.update_transcript(
        transcript_id,
        name=req.name,
        content=req.content,
        language=req.language,
        timeline=req.timeline,
        role_keywords=req.role_keywords,
        labels=req.labels,
    )
    if not updated:
        raise HTTPException(status_code=404, detail="transcript_not_found")
    return updated


@app.delete("/transcripts/{transcript_id}")
def delete_transcript(transcript_id: int):
    transcript = db.get_transcript(transcript_id)
    if not transcript:
        raise HTTPException(status_code=404, detail="transcript_not_found")
    db.delete_transcript(transcript_id)
    return {"deleted": transcript_id}



@app.post("/criteria/parse")
def parse_criteria(
    session_id: str = Form(default=""),
    profile_id: int = Form(default=0),
    google_sheet_id: str = Form(default=""),
    google_gid: str = Form(default=""),
    raw_text: str = Form(default=""),
    file: UploadFile | None = File(default=None),
):
    sid = (session_id or "").strip()
    criteria: list[str] | None = None
    source = None
    try:
        if file is not None:
            content = file.file.read()
            if not content:
                raise ValueError("empty_file")
            suffix = Path(file.filename or "").suffix.lower()
            if suffix == ".csv":
                criteria = _parse_csv(_decode_bytes(content))
                source = "file:csv"
            elif suffix in {".xlsx", ".xlsm"}:
                criteria = _parse_xlsx(content)
                source = "file:xlsx"
            else:
                raise ValueError("unsupported_file")
        elif google_sheet_id.strip():
            sheet_id, gid = _extract_sheet_id(google_sheet_id)
            csv_text = _load_google_sheet_csv(sheet_id, gid or (google_gid.strip() or None))
            criteria = _parse_csv(csv_text)
            source = "google_sheets"
        elif raw_text.strip():
            criteria = _parse_text_block(raw_text)
            source = "text"
        else:
            raise ValueError("no_source")
    except ValueError as exc:
        code = str(exc)
        if code == "unsupported_file":
            raise HTTPException(
                status_code=422,
                detail={"code": "unsupported_file", "message": "Разрешены только CSV и XLSX"},
            ) from exc
        if code == "empty_file":
            raise HTTPException(
                status_code=422,
                detail={"code": "empty_file", "message": "Файл критериев пустой"},
            ) from exc
        if code == "google_sheet_id_required":
            raise HTTPException(
                status_code=422,
                detail={"code": "google_sheet_id_required", "message": "Укажите ID таблицы Google Sheets"},
            ) from exc
        raise HTTPException(
            status_code=400,
            detail={"code": "criteria_source_missing", "message": "Не указан источник критериев"},
        ) from exc
    except RuntimeError as exc:
        if str(exc) == "openpyxl_not_available":
            raise HTTPException(
                status_code=500,
                detail={"code": "openpyxl_missing", "message": "Установите зависимость openpyxl для чтения XLSX"},
            ) from exc
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=502,
            detail={"code": "criteria_fetch_failed", "message": str(exc)},
        ) from exc
    if not criteria:
        raise HTTPException(
            status_code=422,
            detail={"code": "criteria_not_found", "message": "Не удалось выделить критерии"},
        )
    if sid:
        session = db.get_session(sid)
        if not session:
            raise HTTPException(status_code=404, detail={"code": "session_not_found", "message": "Сессия не найдена"})
        db.save_session_criteria(sid, criteria)
        profile_id = session["profile_id"]
    elif profile_id:
        _get_profile_or_404(profile_id)
    return {"criteria": criteria, "count": len(criteria), "session_id": sid or None, "profile_id": profile_id or None, "source": source}

def _build_timeline(segments: list[dict]) -> list[dict]:
    """Объединяет соседние сегменты одного говорящего в таймлайн."""

    timeline: list[dict] = []
    if not segments:
        return timeline

    for segment in segments:
        if not isinstance(segment, dict):
            continue
        text = (segment.get("text") or "").strip()
        if not text:
            continue
        start = float(segment.get("start", 0.0))
        end = float(segment.get("end", start))
        speaker = (segment.get("speaker") or "unknown").strip().lower() or "unknown"

        if timeline and timeline[-1]["speaker"] == speaker and start <= timeline[-1]["end"] + 0.6:
            timeline[-1]["end"] = max(timeline[-1]["end"], end)
            timeline[-1]["text"] = (timeline[-1]["text"] + " " + text).strip()
        else:
            timeline.append(
                {
                    "speaker": speaker,
                    "start": start,
                    "end": end,
                    "text": text,
                }
            )

    return timeline

@app.post("/chat")
def chat(req: ChatReq):
    try:
        session_id = (req.session_id or "").strip() or None
        model_from_req = req.model or req.model_name
        session = None
        profile_id = req.profile_id

        if session_id:
            session = db.get_session(session_id)
            if not session:
                raise HTTPException(status_code=404, detail="session_not_found")
            profile_id = session["profile_id"]
            model_name = session["model_name"]
        else:
            if profile_id is None:
                raise HTTPException(status_code=400, detail="profile_id_required")
            _get_profile_or_404(profile_id)
            model_name = _normalize_model_name(model_from_req)
            title = _session_title_hint(req.title, req.query)
            session = db.create_session(profile_id, title=title, model_name=model_name)
            session_id = session["id"]

        model_name = session["model_name"] if session else _normalize_model_name(model_from_req)
        settings_bundle = _effective_model_settings(profile_id, model_name)  # type: ignore[arg-type]
        effective_settings = settings_bundle.get("effective") or {}
        mem_turns = int(effective_settings.get("mem_turns") or MEM_TURNS)

        history = db.get_messages(session_id, limit_pairs=mem_turns) if session_id else []

        request_criteria = [c for c in (req.criteria or []) if isinstance(c, str)]
        cleaned_criteria = _criteria_from_lines(request_criteria) if request_criteria else []
        if cleaned_criteria and session_id:
            db.save_session_criteria(session_id, cleaned_criteria)
        stored_criteria = db.get_session_criteria(session_id) if session_id else []
        active_criteria = cleaned_criteria or stored_criteria or []

        transcript_text = (req.transcript or "").strip()
        if not transcript_text and session_id:
            latest_transcript = db.get_latest_transcript_for_session(session_id)
            transcript_text = (latest_transcript.get("content") if latest_transcript else "") or ""

        messages = [{"role": "system", "content": _system_prompt()}]
        if req.context:
            messages.append({"role": "user", "content": f"Контекст:\n{req.context}"})
        if transcript_text:
            label = "Транскрипт из запроса" if (req.transcript or "").strip() else "Транскрипт из памяти сессии"
            messages.append({"role": "user", "content": f"{label}:\n{transcript_text}"})
        if active_criteria:
            criteria_block = "\n".join(f"- {item}" for item in active_criteria)
            messages.append({"role": "user", "content": f"Критерии оценки:\n{criteria_block}"})
        # добавляем прошлые реплики (user/assistant) если есть
        if history:
            messages.extend(history)
        # текущий запрос
        user_query = (req.query or "").strip() or "проанализируй"
        messages.append({"role": "user", "content": user_query})

        target_model = _normalize_model_name(model_name)

        resp = client.chat.completions.create(
            model=target_model,
            messages=messages,
            temperature=effective_settings.get("temperature", TEMP),
            max_tokens=effective_settings.get("max_tokens", MAX_TOK),
            stop=["\n\n\n", "\n\n—", "\n\n1."],  # мягкие ограничители списков/абзацев
            presence_penalty=0.2,
            frequency_penalty=0.2,
        )
        text = resp.choices[0].message.content or ""
        text = _sanitize(text)

        if session_id:
            db.add_message(session_id, "user", user_query)
            db.add_message(session_id, "assistant", text)

        return {
            "reply": text,
            "session_id": session_id,
            "profile_id": profile_id,
            "model_name": target_model,
            "used_settings": effective_settings,
        }
    except Exception as e:
        # Возвращаем понятную ошибку, если Ollama/LLM недоступны
        raise HTTPException(status_code=502, detail=f"LLM error: {e}")

LANGUAGE_RE = re.compile(r"^[a-z0-9]{2,8}(?:-[a-z0-9]{2,8})?$")


@app.post("/transcribe")
def transcribe(
    file: UploadFile = File(...),
    model: str = Form(default="small"),
    prompt_preset: str = Form(default="none"),
    language: str = Form(default="auto"),
    session_id: str = Form(default=""),
    profile_id: int = Form(default=0),
):
    """Транскрибация аудио/видео через faster-whisper.

    - model — размер модели whisper (tiny/base/small/medium/large).
    - prompt_preset — режим промпта: none / ru / en.
    - language — `auto` (по умолчанию) или BCP-47 код (`ru`, `en`, `pt-br`...), влияет на второй проход Whisper.
    """
    # сохраняем файл во временный путь
    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=os.path.splitext(file.filename or "")[-1],
    ) as tmp:
        tmp.write(file.file.read())
        tmp_path = tmp.name

    # модель по умолчанию, если в форме пришло пусто
    model_name = (model or os.getenv("WHISPER_MODEL_SIZE", "medium")).strip()

    preset = (prompt_preset or "none").strip().lower()
    if not preset:
        preset = "none"
    if preset not in {"none", "ru", "en"}:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "invalid_prompt_preset",
                "message": "prompt_preset must be one of: none, ru, en.",
            },
        )
    preset_value = None if preset == "none" else preset

    language_raw = (language or "auto").strip().lower()
    if language_raw == "":
        language_raw = "auto"
    if language_raw != "auto" and not LANGUAGE_RE.fullmatch(language_raw):
        raise HTTPException(
            status_code=422,
            detail={
                "code": "invalid_language",
                "message": "language must be 'auto' или BCP-47 код (например, 'ru', 'en', 'pt-br').",
            },
        )
    preferred_language = None if language_raw == "auto" else language_raw

    sid = (session_id or "").strip()
    profile_id_val = profile_id or 0
    if sid:
        session = db.get_session(sid)
        if not session:
            raise HTTPException(status_code=404, detail={"code": "session_not_found", "message": "Сессия не найдена"})
        profile_id_val = session["profile_id"]
    if not profile_id_val:
        profile_id_val = db.ensure_default_profile()
    _get_profile_or_404(profile_id_val)

    try:
        transcription = transcribe_with_fallback(
            audio_path=tmp_path,
            model_name=model_name,
            prompt_preset=preset_value,
            preferred_language=preferred_language,
        )
        text = transcription.get("text", "") if isinstance(transcription, dict) else str(transcription)
        segments = []
        if isinstance(transcription, dict):
            segments = list(transcription.get("segments") or [])
        if not text:
            raise RuntimeError("empty_transcription")
            timeline = _build_timeline(segments)
        detected_language = None
        if isinstance(transcription, dict):
            detected_language = transcription.get("language")
        role_keywords = {}
        if isinstance(transcription, dict):
            role_keywords = transcription.get("role_keywords") or {}

        # Сохраняем транскрипт в БД
        transcript_record = db.create_transcript(
            profile_id=profile_id_val,
            session_id=sid or None,
            name=os.path.basename(file.filename or tmp_path) or "Транскрипт",
            content=text,
            source_path=file.filename,
            language=detected_language or language_raw,
            timeline=timeline,
            role_keywords=role_keywords,
        )

        return {
            "text": text,
            "model": model_name,
            "prompt_preset": preset,
            "language": language_raw,
            "detected_language": detected_language,
            "segments": segments,
            "timeline": timeline,
            "role_keywords": role_keywords,
            "transcript_id": transcript_record.get("id") if transcript_record else None,
            "profile_id": profile_id_val,
            "session_id": sid or None,
        }
    except RuntimeError as e:
        msg = str(e)
        if msg.startswith("WHISPER_MODEL_NOT_FOUND::"):
            missing_model = msg.split("::", 1)[1] if "::" in msg else model_name
            raise HTTPException(
                status_code=424,
                detail={
                    "code": "whisper_model_missing",
                    "model": missing_model,
                    "message": f'Whisper model "{missing_model}" is not available locally.',
                },
            )
        if msg == "empty_transcription":
            raise HTTPException(
                status_code=502,
                detail={
                    "code": "empty_transcription",
                    "message": "Whisper вернул пустую транскрипцию.",
                },
            )
        # остальные ошибки пробрасываем дальше как 500
        raise
    finally:
        try:
            os.remove(tmp_path)
        except Exception:
            pass

@app.get("/health")
def health():
    return {
        "status": "ok",
        "env": ENV,
        "model": MODEL,
        "num_ctx": NUM_CTX,
        "max_tokens": MAX_TOK,
        "mem_turns": MEM_TURNS,
        "temperature": TEMP,
    }


# Lightweight version endpoint to introspect config
# Lightweight version endpoint to introspect config
@app.get("/version")
def version():
    return {"env": ENV, "model": MODEL, "base_url": BASE_URL}


# Reset session endpoint and request model
class ResetReq(BaseModel):
    session_id: str
    admin_token: str | None = None

@app.post("/session/reset")
def reset_session(req: ResetReq):
    if ADMIN_TOKEN and req.admin_token != ADMIN_TOKEN:
        raise HTTPException(status_code=403, detail="forbidden")
    sid = req.session_id.strip()
    db.reset_session(sid)
    return {"ok": True}


from fastapi.middleware.cors import CORSMiddleware

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
    allow_credentials=False,
)

app.mount("/ui", StaticFiles(directory=UI_DIR, html=True), name="ui")

@app.get("/", include_in_schema=False)
def root():
    return RedirectResponse(url="/ui/gui_html.html")
