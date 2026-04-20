from __future__ import annotations

import csv
import io
import re
import urllib.parse
import urllib.request
from typing import Iterable

MAX_CRITERIA = 200


def _normalize_criterion(raw: str) -> str:
    if not raw:
        return ""
    text = re.sub(r"[\u200b\u200c\u200d]", "", raw)
    text = text.strip()
    text = re.sub(r"^[-•*\d\)\].\s]+", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _criteria_from_lines(lines: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for line in lines:
        cleaned = _normalize_criterion(line)
        if not cleaned:
            continue
        key = cleaned.lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(cleaned)
        if len(result) >= MAX_CRITERIA:
            break
    return result


def decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def parse_csv(text: str) -> list[str]:
    stream = io.StringIO(text)
    reader = csv.reader(stream)
    lines: list[str] = []
    for row in reader:
        if not row:
            continue
        for cell in row:
            if cell and cell.strip():
                lines.append(cell)
                break
    return _criteria_from_lines(lines)


def parse_xlsx(data: bytes) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - зависит от окружения
        raise RuntimeError("openpyxl_not_available") from exc

    workbook = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    lines: list[str] = []
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        for cell in row:
            if isinstance(cell, str) and cell.strip():
                lines.append(cell)
                break
            if isinstance(cell, (int, float)):
                lines.append(str(cell))
                break
        if len(lines) >= MAX_CRITERIA:
            break
    return _criteria_from_lines(lines)


def parse_text_block(text: str) -> list[str]:
    if not text:
        return []
    raw_lines = re.split(r"[\r\n]+", text)
    return _criteria_from_lines(raw_lines)


def normalize_criteria_list(values: Iterable[str]) -> list[str]:
    return _criteria_from_lines(values)


def extract_sheet_id(raw: str) -> tuple[str, str | None]:
    value = (raw or "").strip()
    if not value:
        raise ValueError("google_sheet_id_required")
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", value)
    gid = None
    if match:
        sheet_id = match.group(1)
        gid_match = re.search(r"gid=([0-9]+)", value)
        if gid_match:
            gid = gid_match.group(1)
        return sheet_id, gid
    return value, None


def load_google_sheet_csv(sheet_id: str, gid: str | None = None) -> str:
    if not sheet_id:
        raise ValueError("google_sheet_id_required")
    params = {"format": "csv"}
    if gid:
        params["gid"] = gid
    query = urllib.parse.urlencode(params)
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?{query}"
    with urllib.request.urlopen(url, timeout=15) as response:  # type: ignore[call-arg]
        data = response.read()
    return decode_bytes(data)
